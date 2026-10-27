# -*- coding: utf-8 -*-
"""supa_admin.py — 閃耀之星初選平台的管理腳本（在 Scarlett 的電腦執行；需要 Supabase service role 密鑰）

密鑰不放在 repo：在 repo 以外建立一個 .env 檔（例如 ~/ScarlettAIOS/.../工作檔_個資/帳號/supabase.env），內容兩行：
  SUPABASE_URL=https://xxxx.supabase.co
  SUPABASE_SERVICE_ROLE_KEY=（Supabase → Project Settings → API Keys → Secret key，以 sb_secret_ 開頭；或舊式 service_role）

用法：
  python3 -I supa_admin.py <env檔> import   <data_第一版.json>      匯入／更新 229 位學生資料、相片、報告設定（只新增或覆寫，不刪評審紀錄）
  python3 -I supa_admin.py <env檔> accounts <初選平台帳號_v1.xlsx>  按 Excel 建立帳號；已存在的帳號只更新名稱、角色、權重、啟用狀態
  python3 -I supa_admin.py <env檔> weights  <初選平台帳號_v1.xlsx>  只把 Excel 的評分權重及啟用狀態寫入平台
  python3 -I supa_admin.py <env檔> purge                             評審期完結：清除雲端學生資料及相片（保留編號、評審紀錄、委員會決定）
"""
import sys, json, urllib.request, urllib.error

def env(path):
    d = {}
    for line in open(path, encoding="utf-8"):
        line = line.strip()
        if line and not line.startswith("#") and "=" in line:
            k, v = line.split("=", 1); d[k.strip()] = v.strip()
    return d

def call(E, method, path, body=None, headers=None):
    k = E["SUPABASE_SERVICE_ROLE_KEY"]
    h = {"apikey": k, "Content-Type": "application/json"}
    if not k.startswith("sb_"):  # 舊式 JWT 密鑰要同時放 Authorization；新式 sb_secret_ 只放 apikey
        h["Authorization"] = "Bearer " + k
    h.update(headers or {})
    req = urllib.request.Request(E["SUPABASE_URL"].rstrip("/") + path, method=method, headers=h,
                                 data=None if body is None else json.dumps(body, ensure_ascii=False).encode("utf-8"))
    try:
        with urllib.request.urlopen(req, timeout=120) as r:
            t = r.read().decode("utf-8"); return r.status, (json.loads(t) if t else None)
    except urllib.error.HTTPError as e:
        t = e.read().decode("utf-8", "replace"); return e.code, t

def upsert(E, table, rows, conflict):
    for i in range(0, len(rows), 20):
        st, res = call(E, "POST", f"/rest/v1/{table}?on_conflict={conflict}", rows[i:i + 20], {"Prefer": "resolution=merge-duplicates,return=minimal"})
        if st >= 300: raise SystemExit(f"{table} 寫入失敗（{st}）：{res}")

def do_import(E, path):
    D = json.load(open(path, encoding="utf-8"))
    studs, photos = [], []
    for s in D["students"]:
        s = dict(s); ph = s.pop("photo", None)
        studs.append({"id": s["id"], "data": s}); photos.append({"id": s["id"], "photo": ph})
    upsert(E, "students", studs, "id")
    upsert(E, "student_photos", photos, "id")
    upsert(E, "meta", [{"key": "meta", "value": D["meta"]}, {"key": "quality", "value": D.get("quality", [])}], "key")
    print(f"已匯入 {len(studs)} 位學生、{sum(1 for p in photos if p['photo'])} 張相片、報告設定 2 項")

def read_accounts(xlsx):
    import openpyxl
    ws = openpyxl.load_workbook(xlsx, data_only=False)["帳號"]
    head = [c.value for c in ws[4]]; idx = {h: i for i, h in enumerate(head)}
    out = []
    for row in ws.iter_rows(min_row=5, values_only=True):
        u = row[idx["帳號（登入名稱）"]]
        if not u: continue
        out.append({"username": str(u).strip(), "display_name": row[idx["顯示名稱"]], "role": "admin" if "管理員" in str(row[idx["角色"]]) else "judge",
                    "password": str(row[idx["初始密碼"]]), "weight": float(row[idx["評分權重"]] or 1), "active": str(row[idx["狀態"]]).strip() == "啟用"})
    return out

def list_users(E):
    users, page = [], 1
    while True:
        st, res = call(E, "GET", f"/auth/v1/admin/users?page={page}&per_page=100")
        if st >= 300: raise SystemExit(f"讀取帳號失敗（{st}）：{res}")
        batch = res.get("users", []) if isinstance(res, dict) else []
        users += batch
        if len(batch) < 100: return users
        page += 1

def do_accounts(E, xlsx, weights_only=False):
    dom = E.get("EMAIL_DOMAIN", "gef16.invalid")
    existing = {u["email"]: u["id"] for u in list_users(E)}
    profiles = []
    for a in read_accounts(xlsx):
        email = f"{a['username'].lower()}@{dom}"
        uid = existing.get(email)
        if not uid:
            if weights_only: print(f"略過 {a['username']}：平台未有此帳號"); continue
            st, res = call(E, "POST", "/auth/v1/admin/users", {"email": email, "password": a["password"], "email_confirm": True})
            if st >= 300: raise SystemExit(f"建立 {a['username']} 失敗（{st}）：{res}")
            uid = res["id"]; print(f"已建立 {a['username']}（{a['display_name']}）")
        else:
            print(f"已存在 {a['username']}：只更新設定，不改密碼")
        profiles.append({"user_id": uid, "username": a["username"], "display_name": a["display_name"], "role": a["role"], "weight": a["weight"], "active": a["active"]})
    if weights_only:
        for p in profiles:
            call(E, "PATCH", f"/rest/v1/profiles?user_id=eq.{p['user_id']}", {"weight": p["weight"], "active": p["active"]}, {"Prefer": "return=minimal"})
    else:
        upsert(E, "profiles", profiles, "user_id")
    print(f"完成：{len(profiles)} 個帳號")

def do_purge(E):
    ok = input("確定清除雲端學生資料及相片？請輸入「確定清除」：")
    if ok.strip() != "確定清除": print("已取消"); return
    st, res = call(E, "PATCH", "/rest/v1/student_photos?id=neq.__", {"photo": None}, {"Prefer": "return=minimal"})
    st2, res2 = call(E, "GET", "/rest/v1/students?select=id")
    rows = [{"id": r["id"], "data": {"id": r["id"], "purged": True}} for r in (res2 or [])]
    upsert(E, "students", rows, "id")
    print(f"已清除 {len(rows)} 位學生的資料及相片；評審紀錄及委員會決定保留")

if __name__ == "__main__":
    if len(sys.argv) < 3: print(__doc__); sys.exit(1)
    E = env(sys.argv[1]); cmd = sys.argv[2]
    if cmd == "import": do_import(E, sys.argv[3])
    elif cmd == "accounts": do_accounts(E, sys.argv[3])
    elif cmd == "weights": do_accounts(E, sys.argv[3], weights_only=True)
    elif cmd == "purge": do_purge(E)
    else: print(__doc__)
